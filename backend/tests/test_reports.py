"""Reports and their exports, including regressions from the 2026-10-08 audit of the earlier app."""

import csv
import io
import json
import zipfile
from collections.abc import Callable
from typing import Any

import openpyxl
import orjson
import pytest
from fastapi.testclient import TestClient
from pypdf import PdfReader
from sqlalchemy import delete
from sqlalchemy.orm import Session

from app.domain.analytics import SECTIONS
from app.domain.enums import ENUMS
from app.models import DatasetAnalytics, Report
from app.services.reports import content_disposition
from app.utils.format import number, percent

Headers = dict[str, str]
TITLE = "Звіт за Q1/2025 «Огляд»"


@pytest.fixture
def dataset_id(
    client: TestClient,
    auth_headers: Headers,
    run_jobs: Callable[[], int],
    mini_document: dict[str, Any],
) -> int:
    response = client.post(
        "/api/v1/datasets",
        headers=auth_headers,
        files={"file": ("mini.json", orjson.dumps(mini_document), "application/json")},
    )
    run_jobs()
    return response.json()["dataset"]["id"]


def create_report(client: TestClient, headers: Headers, dataset_id: int, **extra: Any) -> dict:
    response = client.post(
        "/api/v1/reports", headers=headers, json={"dataset_id": dataset_id, "title": TITLE, **extra}
    )
    assert response.status_code == 201, response.text
    return response.json()


def export(client: TestClient, headers: Headers, report_id: int, fmt: str) -> Any:
    response = client.get(f"/api/v1/reports/{report_id}/export?format={fmt}", headers=headers)
    assert response.status_code == 200, response.text
    return response


def test_report_lifecycle(client: TestClient, auth_headers: Headers, dataset_id: int) -> None:
    report = create_report(client, auth_headers, dataset_id)
    assert report["title"] == TITLE
    assert report["sections"] == list(SECTIONS)
    assert report["as_of"] == "2026-03-31"
    listed = client.get(f"/api/v1/reports?dataset_id={dataset_id}", headers=auth_headers).json()
    assert [r["id"] for r in listed["items"]] == [report["id"]]
    detail = client.get(f"/api/v1/reports/{report['id']}", headers=auth_headers).json()
    assert [s["key"] for s in detail["content"]] == list(SECTIONS)
    assert detail["content"][0]["kpis"][0] == {
        "key": "members",
        "value": 5,
        "unit": "count",
        "window_days": None,
    }
    assert client.delete(f"/api/v1/reports/{report['id']}", headers=auth_headers).status_code == 204
    assert client.get(f"/api/v1/reports/{report['id']}", headers=auth_headers).status_code == 404


def test_sections_keep_the_canonical_order(
    client: TestClient, auth_headers: Headers, dataset_id: int
) -> None:
    report = create_report(client, auth_headers, dataset_id, sections=["search", "overview"])
    assert report["sections"] == ["overview", "search"]


def test_pdf_has_the_title_every_section_and_cyrillic(
    client: TestClient, auth_headers: Headers, dataset_id: int
) -> None:
    """Audit 2026-10-08: the PDF title was always "Unnamed", 3 of 5 sections were missing and
    Cyrillic letters were not rendered; a title with "/" broke the download."""
    report = create_report(client, auth_headers, dataset_id)
    response = export(client, auth_headers, report["id"], "pdf")
    assert response.headers["content-type"] == "application/pdf"
    disposition = response.headers["content-disposition"]
    assert f'filename="report-{report["id"]}.pdf"' in disposition
    assert "Q1-2025" in disposition
    assert "/" not in disposition.split("filename*=")[1]
    pdf = PdfReader(io.BytesIO(response.content))
    assert pdf.metadata is not None
    assert pdf.metadata.title == TITLE
    text = "\n".join(page.extract_text() for page in pdf.pages)
    assert TITLE in text
    for key in SECTIONS:
        assert ENUMS["analytics_section"][key] in text, key
    assert "Платні підписники" in text
    assert "Тигролови" in text  # a title from the data, in Cyrillic


def test_xlsx_has_a_summary_and_a_sheet_per_chart(
    client: TestClient, auth_headers: Headers, dataset_id: int
) -> None:
    report = create_report(client, auth_headers, dataset_id, sections=["overview", "content"])
    response = export(client, auth_headers, report["id"], "xlsx")
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    summary = workbook["Звіт"]
    assert summary["A1"].value == TITLE
    kpis = {row[1]: row[2] for row in summary.iter_rows(min_row=7, values_only=True) if row[1]}
    assert kpis["Читачі"] == 5
    assert kpis["Відтік"] == 0.5
    assert kpis["Бронювання в черзі"] == 3
    genres = workbook[next(name for name in workbook.sheetnames if "Жанри" in name)]
    assert [cell.value for cell in genres[2]][:3] == ["Жанр", "Позичення", "Читачі"]
    assert genres["A3"].value == "Поезія й драматургія"  # labels, not codes
    assert genres["B3"].value == 2


def test_csv_zip_has_raw_tables(client: TestClient, auth_headers: Headers, dataset_id: int) -> None:
    report = create_report(client, auth_headers, dataset_id)
    response = export(client, auth_headers, report["id"], "csv")
    assert response.headers["content-type"] == "application/zip"
    assert 'filename="report-' in response.headers["content-disposition"]
    archive = zipfile.ZipFile(io.BytesIO(response.content))
    assert "kpis.csv" in archive.namelist()
    assert "usage-reading_heatmap.csv" in archive.namelist()
    raw = archive.read("kpis.csv")
    assert raw.startswith(b"\xef\xbb\xbf")  # BOM: Excel reads it as UTF-8
    kpis = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
    churn = next(r for r in kpis if r["section"] == "overview" and r["metric"] == "churn_rate")
    assert churn == {
        "section": "overview",
        "metric": "churn_rate",
        "label": "Відтік",
        "value": "0.5",
        "unit": "ratio",
        "window_days": "30",
    }
    genres = list(
        csv.DictReader(io.StringIO(archive.read("content-genres.csv").decode("utf-8-sig")))
    )
    assert genres[0]["genre"] == "poetry_drama"  # codes, for machines


def test_json_export(client: TestClient, auth_headers: Headers, dataset_id: int) -> None:
    report = create_report(client, auth_headers, dataset_id, sections=["search"])
    body = json.loads(export(client, auth_headers, report["id"], "json").content)
    assert body["title"] == TITLE
    assert body["as_of"] == "2026-03-31"
    assert [s["key"] for s in body["sections"]] == ["search"]


def test_reports_are_frozen(
    client: TestClient, auth_headers: Headers, dataset_id: int, db_session: Session
) -> None:
    """A report keeps its numbers even when the dataset's analytics change or go away."""
    report = create_report(client, auth_headers, dataset_id)
    db_session.execute(delete(DatasetAnalytics).where(DatasetAnalytics.dataset_id == dataset_id))
    db_session.commit()
    assert export(client, auth_headers, report["id"], "json").status_code == 200


def test_nested_dataset_report_shows_the_notes(
    client: TestClient,
    auth_headers: Headers,
    run_jobs: Callable[[], int],
    nested_document: dict[str, Any],
) -> None:
    response = client.post(
        "/api/v1/datasets",
        headers=auth_headers,
        files={"file": ("nested.json", orjson.dumps(nested_document), "application/json")},
    )
    run_jobs()
    report = create_report(client, auth_headers, response.json()["dataset"]["id"])
    pdf = PdfReader(io.BytesIO(export(client, auth_headers, report["id"], "pdf").content))
    text = " ".join(" ".join(page.extract_text().split()) for page in pdf.pages)
    assert "немає результатів продовження підписок" in text
    for fmt in ("xlsx", "csv", "json"):
        export(client, auth_headers, report["id"], fmt)


def test_report_needs_a_ready_dataset(
    client: TestClient, auth_headers: Headers, mini_document: dict[str, Any]
) -> None:
    pending = client.post(
        "/api/v1/datasets",
        headers=auth_headers,
        files={"file": ("mini.json", orjson.dumps(mini_document), "application/json")},
    ).json()["dataset"]["id"]
    response = client.post(
        "/api/v1/reports", headers=auth_headers, json={"dataset_id": pending, "title": "x"}
    )
    assert response.status_code == 409
    assert response.json()["code"] == "dataset.not_ready"


@pytest.mark.parametrize(
    "body",
    [{"title": ""}, {"title": "x" * 201}, {"sections": []}, {"sections": ["horoscope"]}],
)
def test_invalid_report_requests(
    client: TestClient, auth_headers: Headers, dataset_id: int, body: dict[str, Any]
) -> None:
    response = client.post(
        "/api/v1/reports",
        headers=auth_headers,
        json={"dataset_id": dataset_id, "title": "x", **body},
    )
    assert response.status_code == 422


def test_unknown_export_format(client: TestClient, auth_headers: Headers, dataset_id: int) -> None:
    report = create_report(client, auth_headers, dataset_id)
    response = client.get(
        f"/api/v1/reports/{report['id']}/export?format=docx", headers=auth_headers
    )
    assert response.status_code == 422


def test_reports_are_private_and_follow_their_dataset(
    client: TestClient, auth_headers: Headers, dataset_id: int
) -> None:
    report = create_report(client, auth_headers, dataset_id)
    client.post(
        "/api/v1/auth/register",
        json={"email": "other@example.com", "name": "Інший", "password": "password-2"},
    )
    token = client.post(
        "/api/v1/auth/login", json={"email": "other@example.com", "password": "password-2"}
    ).json()["access_token"]
    other = {"Authorization": f"Bearer {token}"}
    for path in (f"/api/v1/reports/{report['id']}", f"/api/v1/reports/{report['id']}/export"):
        assert client.get(path, headers=other).status_code == 404
    assert client.get("/api/v1/reports", headers=other).json()["total"] == 0
    other_dataset = client.post(
        "/api/v1/reports", headers=other, json={"dataset_id": dataset_id, "title": "x"}
    )
    assert other_dataset.status_code == 404

    assert client.delete(f"/api/v1/datasets/{dataset_id}", headers=auth_headers).status_code == 204
    assert client.get(f"/api/v1/reports/{report['id']}", headers=auth_headers).status_code == 404


@pytest.mark.parametrize(
    ("title", "name"),
    [
        ("Q1/2025", "Q1-2025.pdf"),
        ("../../etc/passwd", "etc-passwd.pdf"),
        (
            "Звіт: вересень?",
            "%D0%97%D0%B2%D1%96%D1%82-%20%D0%B2%D0%B5%D1%80%D0%B5%D1%81%D0%B5%D0%BD%D1%8C.pdf",
        ),
        ("///", "report.pdf"),
    ],
)
def test_file_names_are_safe(title: str, name: str) -> None:
    report = Report(id=7, title=title)
    disposition = content_disposition(report, "pdf")
    assert disposition.startswith('attachment; filename="report-7.pdf"; ')
    assert disposition.endswith(f"filename*=UTF-8''{name}")


def test_ukrainian_number_format() -> None:
    assert number(1234567.891, 2) == "1 234 567,89"
    assert percent(0.0617) == "6,2 %"
